# SPDX-FileCopyrightText: 2025-2026 Taras Paruta (partarstu@gmail.com)
#
# SPDX-License-Identifier: AGPL-3.0-only

"""Attachment extraction, rendering, normalization, and OCR for document RAG."""

import asyncio
import io
import re
import warnings
from dataclasses import dataclass
from typing import Any

import config
from common import utils
from rag_sync import office_conversion
from rag_sync.office_conversion import OfficeConversionError

logger = utils.get_logger("rag_attachment_extraction")

RASTER_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".tif"}
CONVERT_ONLY_EXTENSIONS = {".doc", ".ppt", ".odt", ".odp", ".rtf", ".xls", ".ods"}
CONVERT_WITH_FALLBACK_EXTENSIONS = {".docx", ".pptx"}
TEXT_ONLY_CONVERTED_EXTENSIONS = {".xls", ".ods"}
OFFICE_EXTENSIONS = CONVERT_ONLY_EXTENSIONS | CONVERT_WITH_FALLBACK_EXTENSIONS


class UnsupportedFormatError(ValueError):
    """The attachment format is not supported by document ingestion."""


class ExtractionError(RuntimeError):
    """The attachment could not be safely extracted."""


@dataclass(slots=True)
class PageContent:
    """One extracted page with normalized text and an optional PNG image."""

    text: str
    image: bytes | None


@dataclass(slots=True)
class ExtractedDocument:
    """Extracted pages and the source document's true page count."""

    pages: list[PageContent]
    total_page_count: int


def skip_reason(file_name: str) -> str | None:
    """Returns why an attachment can't be ingested at all, or ``None`` when it can.

    Checked before download, so a format that can never be extracted is skipped
    instead of failing (and being retried) on every run.
    """
    extension = _extension_of(file_name)
    supported = RASTER_IMAGE_EXTENSIONS | OFFICE_EXTENSIONS | {".pdf", ".xlsx", ".csv", ".txt", ".md"}
    if extension not in supported:
        return f"attachments of type '{extension or file_name}' are not ingested"
    if extension in CONVERT_ONLY_EXTENSIONS and not office_conversion.conversion_available():
        return f"'{extension}' needs office conversion, which is disabled or unavailable"
    return None


def extract_attachment(file_name: str, content: bytes) -> ExtractedDocument:
    """Extracts a non-office attachment without blocking on an external process."""
    extension = _extension_of(file_name)
    match extension:
        case ".pdf":
            return _extract_pdf(content)
        case ".xlsx":
            return _extract_xlsx(content)
        case ".txt" | ".md":
            return ExtractedDocument([PageContent(_decode_text(content), None)], 1)
        case ".csv":
            return ExtractedDocument([PageContent(_normalize_extracted_text(_decode_text(content)), None)], 1)
        case ext if ext in RASTER_IMAGE_EXTENSIONS:
            return _extract_raster_images(content)
        case ext if ext in OFFICE_EXTENSIONS:
            raise ExtractionError("Office attachments must be extracted with extract_attachment_async().")
        case _:
            raise UnsupportedFormatError(f"Attachments of type '{extension or file_name}' are not ingested.")


async def extract_attachment_async(file_name: str, content: bytes) -> ExtractedDocument:
    """Extracts an attachment without blocking the event loop."""
    extension = _extension_of(file_name)
    if extension in OFFICE_EXTENSIONS:
        return await _extract_office(file_name, content, extension)
    return await asyncio.to_thread(extract_attachment, file_name, content)


def _extract_pdf(content: bytes, include_images: bool = True) -> ExtractedDocument:
    """Extracts native text and, when enabled, a PNG rendering for each PDF page."""
    import pymupdf

    try:
        document = pymupdf.open(stream=content, filetype="pdf")
    except Exception as error:
        raise ExtractionError(f"The PDF could not be opened: {error}") from error

    try:
        with document:
            total_page_count = document.page_count
            page_limit = config.DocumentRagConfig.MAX_PAGES_PER_DOCUMENT
            _log_truncation("PDF", total_page_count, page_limit)
            pages: list[PageContent] = []
            for page_number in range(min(total_page_count, page_limit)):
                page = document[page_number]
                native_text = page.get_text()
                image = _render_page(page) if include_images else None
                text = _page_text_with_ocr(page, native_text, image) if include_images else native_text
                pages.append(PageContent(_normalize_extracted_text(text), image))
            return ExtractedDocument(pages, total_page_count)
    except ExtractionError:
        raise
    except Exception as error:
        raise ExtractionError(f"PDF extraction failed: {error}") from error


def _page_text_with_ocr(page: Any, native_text: str, page_image: bytes | None) -> str:
    """Applies full-page OCR to image-only pages and embedded-image OCR otherwise."""
    from rag_sync import ocr as ocr_module

    if page_image is None:
        return native_text
    threshold = config.DocumentRagConfig.OCR_TEXT_THRESHOLD_CHARACTERS
    if len(native_text.strip()) < threshold:
        return ocr_module.extract_text(page_image) or native_text

    embedded_text: list[str] = []
    seen_xrefs: set[int] = set()
    try:
        for image_info in page.get_images(full=True):
            xref = image_info[0]
            if xref in seen_xrefs:
                continue
            seen_xrefs.add(xref)
            extracted = page.parent.extract_image(xref)
            normalized = _normalize_png(extracted["image"])
            if normalized is not None:
                text = ocr_module.extract_text(normalized)
                if text:
                    embedded_text.append(text)
    except Exception:
        logger.warning("Embedded-image OCR failed on one page; native text is kept.", exc_info=True)
    if not embedded_text:
        return native_text
    return f"{native_text.rstrip()}\n\n[embedded image text]\n{'\n'.join(embedded_text)}"


def _render_page(page: Any) -> bytes | None:
    """Renders one PDF page to a normalized PNG within the configured pixel cap.

    The PIL image is built directly from the pixmap samples (encoding the pixmap to
    PNG and re-decoding it with Pillow would do the pixel work twice).
    """
    from PIL import Image

    try:
        pixmap = page.get_pixmap(dpi=config.DocumentRagConfig.RENDER_DPI)
        max_pixels = config.DocumentRagConfig.MAX_IMAGE_PIXELS
        # Decompression-bomb guard on the raw raster, mirroring the decoder-side cap.
        if pixmap.width * pixmap.height > max_pixels * max_pixels:
            logger.warning(
                "Rendered page exceeds the pixel cap (%dx%d); the page stays text-only.",
                pixmap.width,
                pixmap.height,
            )
            return None
        mode = "RGBA" if pixmap.alpha else "RGB"
        image = Image.frombytes(mode, (pixmap.width, pixmap.height), pixmap.samples)
        return _image_to_png(image)
    except Exception:
        logger.warning("Page rendering failed; the page stays text-only.", exc_info=True)
        return None


def _normalize_png(content: bytes) -> bytes | None:
    """Decodes an image with Pillow's bomb guard, caps its dimensions, and returns PNG."""
    from PIL import Image

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(content)) as image:
                image.load()
                return _image_to_png(image)
    except Exception:
        logger.warning("A page image could not be decoded safely; the page stays text-only.", exc_info=True)
        return None


def _image_to_png(image: Any) -> bytes:
    """Normalizes one loaded Pillow image to RGB PNG within the dimension cap."""
    from PIL import ImageOps

    normalized = ImageOps.exif_transpose(image).convert("RGB")
    max_dimension = config.DocumentRagConfig.MAX_IMAGE_PIXELS
    if max(normalized.size) > max_dimension:
        normalized.thumbnail((max_dimension, max_dimension))
    buffer = io.BytesIO()
    normalized.save(buffer, format="PNG")
    return buffer.getvalue()


def _extract_raster_images(content: bytes) -> ExtractedDocument:
    """Extracts each frame of a raster image as one normalized, OCR'd page."""
    from PIL import Image

    from rag_sync import ocr as ocr_module

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(content)) as image:
                total_page_count = getattr(image, "n_frames", 1)
                page_limit = config.DocumentRagConfig.MAX_PAGES_PER_DOCUMENT
                _log_truncation("Raster image", total_page_count, page_limit)
                pages: list[PageContent] = []
                for page_number in range(min(total_page_count, page_limit)):
                    image.seek(page_number)
                    image.load()
                    normalized = _image_to_png(image)
                    text = _normalize_extracted_text(ocr_module.extract_text(normalized))
                    pages.append(PageContent(text, normalized))
                return ExtractedDocument(pages, total_page_count)
    except Exception as error:
        raise ExtractionError(f"The image could not be decoded safely: {error}") from error


async def _extract_office(file_name: str, content: bytes, extension: str) -> ExtractedDocument:
    """Converts an office document to PDF, with native fallback for DOCX and PPTX."""
    if not office_conversion.conversion_available():
        if extension in CONVERT_WITH_FALLBACK_EXTENSIONS:
            return await asyncio.to_thread(_extract_office_native, extension, content)
        raise OfficeConversionError(f"Conversion of '{file_name}' is unavailable; the format needs it.")

    try:
        pdf_bytes = await office_conversion.convert_to_pdf(content, file_name)
    except OfficeConversionError:
        if extension not in CONVERT_WITH_FALLBACK_EXTENSIONS:
            raise
        logger.warning(f"Converting '{file_name}' failed; falling back to its native text-only reader.")
        return await asyncio.to_thread(_extract_office_native, extension, content)

    include_images = extension not in TEXT_ONLY_CONVERTED_EXTENSIONS
    return await asyncio.to_thread(_extract_pdf, pdf_bytes, include_images)


def _extract_office_native(extension: str, content: bytes) -> ExtractedDocument:
    """Uses text-only native readers when DOCX or PPTX conversion is unavailable."""
    match extension:
        case ".docx":
            from docx import Document

            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
            return ExtractedDocument([PageContent(_normalize_extracted_text(text), None)], 1)
        case ".pptx":
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(content))
            total_page_count = len(presentation.slides)
            page_limit = config.DocumentRagConfig.MAX_PAGES_PER_DOCUMENT
            _log_truncation("PowerPoint", total_page_count, page_limit)
            pages: list[PageContent] = []
            for slide_number, slide in enumerate(presentation.slides):
                if slide_number >= page_limit:
                    break
                text = "\n".join(
                    shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
                )
                pages.append(PageContent(_normalize_extracted_text(text), None))
            return ExtractedDocument(pages, total_page_count)
        case _:
            raise UnsupportedFormatError(f"No native reader for '{extension}'.")


def _extract_xlsx(content: bytes) -> ExtractedDocument:
    """Reads an XLSX workbook sheet by sheet, without conversion or page images."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        total_page_count = len(workbook.worksheets)
        page_limit = config.DocumentRagConfig.MAX_PAGES_PER_DOCUMENT
        _log_truncation("Spreadsheet", total_page_count, page_limit)
        pages: list[PageContent] = []
        for sheet in workbook.worksheets[:page_limit]:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    lines.append("\t".join(cells))
            pages.append(PageContent(_normalize_extracted_text("\n".join(lines)), None))
        return ExtractedDocument(pages, total_page_count)
    finally:
        workbook.close()


def _normalize_extracted_text(text: str) -> str:
    """Normalizes extracted text while retaining paragraph boundaries and content order."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(?<=\w)-\n(?=\w)", "", text)
    lines = []
    for line in text.split("\n"):
        cleaned = re.sub(r"[ \t]+", " ", line).strip()
        if re.fullmatch(r"\d+", cleaned):
            continue
        lines.append(cleaned)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _decode_text(content: bytes) -> str:
    """Decodes text with BOM awareness and a deterministic single-byte fallback."""
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extension_of(file_name: str) -> str:
    """Returns a lower-case filename extension without trusting path components."""
    name = file_name.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    if "." not in name:
        return ""
    return name[name.rfind(".") :].lower()


def _log_truncation(kind: str, total_page_count: int, page_limit: int) -> None:
    if total_page_count > page_limit:
        logger.warning(f"{kind} has {total_page_count} pages; ingesting only the first {page_limit}.")
